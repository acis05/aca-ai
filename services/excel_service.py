from __future__ import annotations
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any
from openpyxl import load_workbook


def _cell_value(v):
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    if isinstance(v, date): return v.isoformat()
    return v


def inspect_workbook(path: str | Path, sample_rows: int = 20):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row: raise ValueError('Excel tidak memiliki header.')
    headers = [str(x).strip() if x is not None else '' for x in header_row]
    while headers and not headers[-1]: headers.pop()
    if not any(headers): raise ValueError('Baris pertama Excel kosong. Gunakan baris pertama sebagai header.')
    nonblank=[h for h in headers if h]
    if len(set(nonblank)) != len(nonblank): raise ValueError('Nama kolom harus unik.')
    samples=[]
    for idx,row in enumerate(rows):
        if idx>=sample_rows: break
        rec={headers[i]:_cell_value(row[i]) for i in range(min(len(headers),len(row))) if headers[i]}
        if any(v not in (None,'') for v in rec.values()): samples.append(rec)
    sheet=ws.title; wb.close()
    return headers,samples,sheet


def inspect_source(path, sample_rows=20): return inspect_workbook(path,sample_rows)


def build_dynamic_schema(target_path: str | Path, sample_rows: int = 5):
    headers,samples,sheet=inspect_workbook(target_path,sample_rows)
    return {
        'platform':'ACA Desktop',
        'module':'Dynamic Excel Mapping',
        'sheet_name':sheet,
        'template_file':str(target_path),
        'target_columns':[{'name':h,'description':_describe_target(h,samples)} for h in headers if h],
        'required_targets':[],
        'normalizers':_infer_normalizers(headers),
        'target_samples':samples,
    }


def _describe_target(header, samples):
    vals=[]
    for r in samples:
        v=r.get(header)
        if v not in (None,''):
            vals.append(str(v)[:60])
        if len(vals)>=2: break
    return ('Contoh target: '+ ' | '.join(vals)) if vals else 'Kolom target dari template ACA yang di-upload.'


def _infer_normalizers(headers):
    out={}
    for h in headers:
        u=h.upper()
        if 'DATE' in u or u.startswith('TGL') or 'TANGGAL' in u: out[h]='date'
    return out


def read_source_rows(path: str | Path):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True); header=next(it,None)
    if not header: wb.close(); return [],[]
    headers=[str(x).strip() if x is not None else '' for x in header]
    while headers and not headers[-1]: headers.pop()
    rows=[]
    for row in it:
        if all(v is None or str(v).strip()=='' for v in row): continue
        rows.append({headers[i]:row[i] for i in range(min(len(headers),len(row))) if headers[i]})
    wb.close(); return headers,rows


def _to_date(value):
    if value in (None,''): return None
    if isinstance(value,datetime): return value
    if isinstance(value,date): return datetime.combine(value,datetime.min.time())
    text=str(value).strip()
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y','%m/%d/%Y','%Y/%m/%d'):
        try:return datetime.strptime(text,fmt)
        except ValueError: pass
    return value


def transform_value(value: Any, rule: dict, target: str, schema: dict):
    if schema.get('normalizers',{}).get(target)=='date': value=_to_date(value)
    return value


def convert_rows(source_rows:list[dict], mapping:dict, schema:dict):
    targets=[x['name'] for x in schema['target_columns']]; rules=mapping.get('target_rules',{}); output=[]
    for row in source_rows:
        out={t:None for t in targets}
        for target,rule in rules.items():
            if target not in out: continue
            mode=rule.get('mode','direct')
            if mode=='constant': raw=rule.get('constant')
            else: raw=row.get(rule.get('source')) if rule.get('source') else None
            out[target]=transform_value(raw,rule,target,schema)
        output.append(out)
    return output


def validate_output(rows:list[dict], schema:dict):
    issues=[]
    mapped=set()
    for row in rows:
        for k,v in row.items():
            if v not in (None,''): mapped.add(k)
    if not mapped: issues.append('Belum ada kolom target yang terisi. Periksa mapping.')
    if rows and all(all(v in (None,'') for v in r.values()) for r in rows): issues.append('Semua hasil mapping kosong.')
    return issues


def write_target(template_path: str | Path, output_path: str | Path, rows:list[dict], schema:dict):
    wb=load_workbook(template_path)
    ws=wb[schema['sheet_name']]
    headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
    while headers and headers[-1] in (None,''): headers.pop()
    style_snapshot=[]
    for c_idx in range(1,len(headers)+1):
        src=ws.cell(2,c_idx) if ws.max_row>=2 else ws.cell(1,c_idx)
        style_snapshot.append({'_style':copy(src._style),'number_format':src.number_format,'alignment':copy(src.alignment),'fill':copy(src.fill),'font':copy(src.font),'border':copy(src.border),'protection':copy(src.protection)})
    if ws.max_row>1: ws.delete_rows(2,ws.max_row-1)
    for r_idx,row in enumerate(rows,start=2):
        for c_idx,header in enumerate(headers,start=1):
            cell=ws.cell(r_idx,c_idx,row.get(header))
            st=style_snapshot[c_idx-1]; cell._style=copy(st['_style']); cell.number_format=st['number_format']; cell.alignment=copy(st['alignment']); cell.fill=copy(st['fill']); cell.font=copy(st['font']); cell.border=copy(st['border']); cell.protection=copy(st['protection'])
    wb.save(output_path)
